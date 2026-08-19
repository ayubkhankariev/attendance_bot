"""Учёт штрафов (нарушений): замена преподавателя, мало студентов, отменённая пара.

Хранится в violations.json (плоский список записей) рядом с config.yaml —
переживает перезапуск бота, как и lang_prefs.json.

TALABALARSONI и DARSBARBOD бот проставляет автоматически при отправке фото
в группу руководителей (см. _maybe_log_violation в main.py) — их можно
определить по счётчику людей на фото. ZAMENA (замена преподавателя другим)
камера подтвердить не может — такие записи добавляются вручную через меню
«Штрафы» → «Добавить нарушение».
"""
from __future__ import annotations

import datetime as dt
import json
import logging
import pathlib
import uuid

log = logging.getLogger("monitor.penalties")

VIOLATIONS_FILE = pathlib.Path("violations.json")

TYPES = ("ZAMENA", "TALABALARSONI", "DARSBARBOD")


def load() -> list[dict]:
    if VIOLATIONS_FILE.exists():
        try:
            return json.loads(VIOLATIONS_FILE.read_text(encoding="utf-8"))
        except Exception:
            log.exception("Не удалось прочитать %s", VIOLATIONS_FILE)
    return []


def save(records: list[dict]) -> None:
    try:
        VIOLATIONS_FILE.write_text(
            json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except Exception:
        log.exception("Не удалось сохранить %s", VIOLATIONS_FILE)


def all_records(app) -> list[dict]:
    return app.bot_data.setdefault("violations", load())


def find_duplicate(records: list[dict], date: dt.date, period: int, 
                   teacher: str, groups: list[str]) -> dict | None:
    """Поиск дубликата по дате, паре, преподавателю и группам.
    Возвращает существующую запись или None."""
    date_str = date.isoformat()
    groups_set = set(groups) if groups else set()
    for r in records:
        if (r["date"] == date_str and r["period"] == period and 
            r["teacher"] == teacher and set(r.get("groups", [])) == groups_set):
            return r
    return None


def add(app, *, date: dt.date, period: int, teacher: str, subject: str,
        groups: list[str], vtype: str, students: int | None = None,
        total: int | None = None, link: str | None = None, 
        message_id: int | None = None, message_chat_id: int | None = None) -> dict:
    records = all_records(app)
    rec = {
        "id": uuid.uuid4().hex[:10],
        "date": date.isoformat(),
        "period": period,
        "teacher": teacher,
        "subject": subject,
        "groups": groups or [],
        "type": vtype,
        "students": students,
        "total": total,
        "link": link,
        "message_id": message_id,
        "message_chat_id": message_chat_id,
        "created": dt.datetime.now().isoformat(timespec="seconds"),
    }
    records.append(rec)
    save(records)
    return rec


def delete(app, record_id: str) -> bool:
    """Удаление записи по ID. Возвращает True если запись найдена и удалена."""
    records = all_records(app)
    for i, r in enumerate(records):
        if r["id"] == record_id:
            records.pop(i)
            save(records)
            return True
    return False


def find_by_message_id(records: list[dict], message_id: int) -> dict | None:
    """Поиск штрафа по ID сообщения в группе."""
    for r in records:
        if r.get("message_id") == message_id:
            return r
    return None


def find_duplicate_within_time(records: list[dict], date: dt.date, period: int, 
                               teacher: str, groups: list[str], 
                               minutes_threshold: int = 5) -> dict | None:
    """Поиск дубликата в течение N минут (по умолчанию 5).
    Возвращает существующую запись или None."""
    date_str = date.isoformat()
    groups_set = set(groups) if groups else set()
    now = dt.datetime.now().isoformat(timespec="seconds")
    threshold_time = (dt.datetime.fromisoformat(now) - 
                     dt.timedelta(minutes=minutes_threshold))
    
    for r in records:
        if (r["date"] == date_str and r["period"] == period and 
            r["teacher"] == teacher and set(r.get("groups", [])) == groups_set):
            # Проверяем, был ли пост отправлен в течение последних N минут
            created = dt.datetime.fromisoformat(r["created"])
            if created > threshold_time:
                return r
    return None


def update(app, record_id: str, **kwargs) -> bool:
    """Обновление полей записи по ID. Возвращает True если обновлено."""
    records = all_records(app)
    for r in records:
        if r["id"] == record_id:
            r.update(kwargs)
            save(records)
            return True
    return False


def update_by_message_id(app, message_id: int, **kwargs) -> bool:
    """Обновление записи по ID сообщения в группе."""
    records = all_records(app)
    for r in records:
        if r.get("message_id") == message_id:
            r.update(kwargs)
            save(records)
            return True
    return False


def delete_by_message_id(app, message_id: int) -> bool:
    """Удаление записи по ID сообщения в группе."""
    records = all_records(app)
    for i, r in enumerate(records):
        if r.get("message_id") == message_id:
            records.pop(i)
            save(records)
            return True
    return False


def filter_records(records: list[dict], date_from: dt.date | None = None,
                    date_to: dt.date | None = None, teacher_query: str | None = None,
                    vtype: str | None = None) -> list[dict]:
    out = []
    for r in records:
        d = dt.date.fromisoformat(r["date"])
        if date_from and d < date_from:
            continue
        if date_to and d > date_to:
            continue
        if vtype and r["type"] != vtype:
            continue
        if teacher_query and teacher_query.strip().upper() not in r["teacher"].upper():
            continue
        out.append(r)
    out.sort(key=lambda r: (r["date"], r["period"]))
    return out


def hashtag_label(rec: dict) -> str:
    if rec["type"] == "TALABALARSONI" and rec.get("students") is not None:
        total = rec.get("total") or "?"
        return f"#TALABALARSONI {rec['students']}/{total}"
    return f"#{rec['type']}"


def export_excel(records: list[dict], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Штрафы"
    headers = ["№", "SANA", "PARASI", "F.ISH.", "FAN VA GURUXLAR", "NARUSHENIYA TURI", "ILOVA"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for i, r in enumerate(records, start=1):
        d = dt.date.fromisoformat(r["date"])
        groups = " / ".join(r.get("groups") or []) or "-"
        subject_groups = f"{r['subject']} {groups}".strip()
        row = i + 1
        ws.append([i, d.strftime("%d/%m/%y"), r["period"], r["teacher"],
                   subject_groups, hashtag_label(r), "TELEG" if r.get("link") else "-"])
        if r.get("link"):
            cell = ws.cell(row=row, column=7)
            cell.hyperlink = r["link"]
            cell.style = "Hyperlink"

    widths = [5, 10, 8, 24, 32, 24, 10]
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
